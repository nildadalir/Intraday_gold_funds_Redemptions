"""BI Excel reader — instruments grouped later by AssetId."""

from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

import config

logger = logging.getLogger(__name__)

REQUIRED_COLUMNS = (
    "InstrumentId",
    "Instrument",
    "AssetId",
    "Asset",
    "TseId",
)


class ExcelReadError(Exception):
    """Excel missing, unreadable, or missing required columns."""


@dataclass(frozen=True)
class InstrumentRow:
    instrument_id: str
    instrument: str
    asset_id: str
    asset: str
    tse_id: str | None
    market: str = ""


def _cell_str(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text or text.upper() == "NULL":
        return None
    return text


def parse_tse_id(value: Any) -> str | None:
    """
    Parse TseId as a digit string (preserve full precision).

    New BI exports store TseId as text; legacy floats are still accepted.
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return str(value) if value > 0 else None
    if isinstance(value, float):
        if value != value:  # NaN
            return None
        return str(int(value))
    text = _cell_str(value)
    if text is None:
        return None
    if text.isdigit():
        return text
    try:
        return str(int(float(text)))
    except ValueError as exc:
        raise ExcelReadError(f"Unparseable TseId: {value!r}") from exc


def load_instrument_rows(excel_path: Path | None = None) -> list[InstrumentRow]:
    path = excel_path or config.EXCEL_PATH
    if not path.exists():
        raise ExcelReadError(f"Excel not found: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header = next(rows_iter, None)
        if not header:
            raise ExcelReadError("Excel has no header row")
        col_index = {
            str(name).strip(): i
            for i, name in enumerate(header)
            if name is not None and str(name).strip()
        }
        missing = [c for c in REQUIRED_COLUMNS if c not in col_index]
        if missing:
            raise ExcelReadError(f"Excel missing columns: {missing}")

        out: list[InstrumentRow] = []
        for raw in rows_iter:
            if raw is None:
                continue

            def col(name: str) -> Any:
                return raw[col_index[name]] if col_index[name] < len(raw) else None

            instrument = _cell_str(col("Instrument"))
            asset = _cell_str(col("Asset"))
            asset_id = _cell_str(col("AssetId"))
            instrument_id = _cell_str(col("InstrumentId"))
            if not instrument or not asset or not asset_id or not instrument_id:
                continue

            market = _cell_str(col("Market")) or ""

            out.append(
                InstrumentRow(
                    instrument_id=instrument_id,
                    instrument=instrument,
                    asset_id=asset_id,
                    asset=asset,
                    tse_id=parse_tse_id(col("TseId")),
                    market=market.strip(),
                )
            )
        logger.info("Loaded %s instrument rows from %s", len(out), path.name)
        return out
    finally:
        wb.close()
